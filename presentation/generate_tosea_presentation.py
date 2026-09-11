from pathlib import Path
from math import cos, sin, pi
import argparse
from typing import Optional

from PIL import Image, ImageDraw, ImageFilter, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUTPUT = ROOT / "泰卦外贸团队系统.pptx"
BG_PATH = ASSETS / "tosea_deep_blue_bg.png"
BG_SIZE = (1600, 900)
GRID_STEP = 80
ARC_SPECS = ((560, 90), (700, 45), (860, 20))
TOP_LEFT_POLYGON = [(0, 0), (500, 0), (260, 240)]
BOTTOM_RIGHT_POLYGON = [(1600, 900), (1180, 900), (1600, 640)]
TOSEA_FRAME = (1000, 90, 1470, 170)
GLOW_ELLIPSES = (
    ((1180, 40, 1510, 370), (26, 148, 255, 90)),
    ((40, 620, 420, 980), (19, 109, 215, 70)),
)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

WHITE = RGBColor(255, 255, 255)
TEXT = RGBColor(229, 240, 255)
MUTED = RGBColor(166, 203, 242)
ACCENT = RGBColor(32, 195, 255)
ACCENT_2 = RGBColor(111, 255, 233)
CARD_FILL = RGBColor(12, 40, 89)
CARD_LINE = RGBColor(71, 145, 234)


def emu(inches: float) -> int:
    return Inches(inches)


def ensure_background(bg_path: Path = BG_PATH) -> None:
    bg_path.parent.mkdir(parents=True, exist_ok=True)
    if bg_path.exists():
        return
    width, height = BG_SIZE
    vertical = ImageOps.colorize(
        Image.linear_gradient("L").resize((width, height)),
        black="#081a39",
        white="#144f97",
    )
    horizontal = ImageOps.colorize(
        Image.linear_gradient("L").rotate(90, expand=True).resize((width, height)),
        black="#081a39",
        white="#0d325e",
    )
    img = Image.blend(vertical, horizontal, 0.42)

    overlay = Image.new("RGBA", (width, height), (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay)

    for step in range(0, width, GRID_STEP):
        draw.line((step, 0, step, height), fill=(64, 128, 255, 20), width=1)
    for step in range(0, height, GRID_STEP):
        draw.line((0, step, width, step), fill=(64, 128, 255, 18), width=1)

    for radius, alpha in ARC_SPECS:
        draw.arc((width - radius - 120, -140, width + radius - 120, height + 140),
                 start=200, end=342, fill=(59, 171, 255, alpha), width=3)

    draw.polygon(TOP_LEFT_POLYGON, fill=(12, 91, 190, 30))
    draw.polygon(BOTTOM_RIGHT_POLYGON, fill=(14, 72, 164, 50))
    draw.rounded_rectangle(TOSEA_FRAME, radius=18, outline=(111, 255, 233, 120), width=2)
    draw.text((1045, 114), "TOSEA", fill=(180, 238, 255, 120))

    glow = Image.new("RGBA", (width, height), (0, 0, 0, 0))
    gdraw = ImageDraw.Draw(glow)
    for bounds, fill in GLOW_ELLIPSES:
        gdraw.ellipse(bounds, fill=fill)
    glow = glow.filter(ImageFilter.GaussianBlur(40))

    img = Image.alpha_composite(img.convert("RGBA"), glow)
    img = Image.alpha_composite(img, overlay)
    img.convert("RGB").save(bg_path)



def set_bg(slide, bg_path: Path = BG_PATH):
    ensure_background(bg_path)
    slide.shapes.add_picture(str(bg_path), 0, 0, width=SLIDE_W, height=SLIDE_H)
    top_bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, SLIDE_W, emu(0.14))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = ACCENT
    top_bar.fill.transparency = 0.15
    top_bar.line.fill.background()


def add_footer(slide, right_text="TOSEA 泰卦外贸"):
    box = slide.shapes.add_textbox(emu(10.7), emu(7.0), emu(2.0), emu(0.25))
    p = box.text_frame.paragraphs[0]
    p.text = right_text
    p.font.size = Pt(10)
    p.font.color.rgb = MUTED
    p.alignment = PP_ALIGN.RIGHT


def set_shape_text(shape, text, size, color=WHITE, bold=False, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE


def add_title(slide, title, subtitle=None, section=None):
    if section is not None:
        tag = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, emu(0.72), emu(0.65), emu(0.95), emu(0.38))
        tag.fill.solid()
        tag.fill.fore_color.rgb = ACCENT
        tag.line.fill.background()
        tf = tag.text_frame
        tf.text = section
        tf.paragraphs[0].font.size = Pt(15)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = RGBColor(7, 29, 67)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    box = slide.shapes.add_textbox(emu(0.9 if section is None else 1.9), emu(0.48), emu(8.8), emu(0.62))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = WHITE
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(11)
        p2.font.color.rgb = MUTED
        p2.space_before = Pt(3)


def add_bullet_list(slide, items, x, y, w, h, font_size=20, accent=False):
    box = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = TEXT
        p.font.bold = accent
        p.space_after = Pt(8)
        p.level = 0
        p.bullet = True
    return box


def add_card(slide, x, y, w, h, title, lines, title_size=18):
    shadow = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, emu(x + 0.05), emu(y + 0.06), emu(w), emu(h))
    shadow.fill.solid()
    shadow.fill.fore_color.rgb = RGBColor(0, 0, 0)
    shadow.fill.transparency = 0.82
    shadow.line.fill.background()
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, emu(x), emu(y), emu(w), emu(h))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_FILL
    card.fill.transparency = 0.12
    card.line.color.rgb = CARD_LINE
    card.line.width = Pt(1.5)

    title_box = slide.shapes.add_textbox(emu(x + 0.22), emu(y + 0.18), emu(w - 0.44), emu(0.32))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(title_size)
    p.font.bold = True
    p.font.color.rgb = WHITE

    body = slide.shapes.add_textbox(emu(x + 0.22), emu(y + 0.58), emu(w - 0.44), emu(h - 0.72))
    tf = body.text_frame
    tf.word_wrap = True
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT
        p.space_after = Pt(6)
    return card


def slide_cover(prs, bg_path: Path = BG_PATH):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg_path)
    ribbon = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, emu(0.95), emu(0.76), emu(1.7), emu(0.34))
    ribbon.fill.solid()
    ribbon.fill.fore_color.rgb = RGBColor(10, 57, 115)
    ribbon.line.color.rgb = ACCENT_2
    set_shape_text(ribbon, "TEAM SYSTEM", 12, ACCENT_2, True)

    title = slide.shapes.add_textbox(emu(0.95), emu(1.7), emu(8.7), emu(1.6))
    tf = title.text_frame
    p = tf.paragraphs[0]
    p.text = "泰卦外贸团队系统"
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p2 = tf.add_paragraph()
    p2.text = "外贸团队规范化管理与业务协同体系"
    p2.font.size = Pt(16)
    p2.font.color.rgb = MUTED
    p2.space_before = Pt(10)

    summary = slide.shapes.add_textbox(emu(0.98), emu(3.45), emu(4.7), emu(1.05))
    tf = summary.text_frame
    for idx, line in enumerate(["聚焦组织协同", "统一制度标准", "强化结果导向"]):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(18)
        p.font.color.rgb = TEXT
        p.space_after = Pt(8)

    panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, emu(8.55), emu(1.55), emu(3.78), emu(4.3))
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(10, 31, 70)
    panel.fill.transparency = 0.18
    panel.line.color.rgb = CARD_LINE
    panel.line.width = Pt(1.5)

    box = slide.shapes.add_textbox(emu(8.9), emu(1.98), emu(3.1), emu(3.4))
    tf = box.text_frame
    items = [
        ("主题", "六大板块管理体系"),
        ("风格", "深蓝 TOSEA 商务科技风"),
        ("输出", "9 页演示文稿，可直接下载使用"),
    ]
    for idx, (k, v) in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"{k}｜{v}"
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT
        p.space_after = Pt(16)
    add_footer(slide)


def slide_overview(prs, bg_path: Path = BG_PATH):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg_path)
    add_title(slide, "核心管理板块", "围绕组织、制度、职责、流程与激励，构建外贸团队协同闭环")
    cards = [
        ("01 团队", ["明确组织结构", "形成高效协作关系"]),
        ("02 团队制度", ["统一执行标准", "保障团队稳定运行"]),
        ("03 岗位职责", ["责任清晰到岗", "协同衔接顺畅"]),
        ("04 工作流程", ["规范日常动作", "提升任务执行效率"]),
        ("05 业务流程", ["打通客户到成交链路", "增强业务可控性"]),
        ("06 薪酬绩效", ["激励与贡献匹配", "推动结果持续增长"]),
    ]
    positions = [(0.95, 1.75), (4.47, 1.75), (7.99, 1.75), (0.95, 4.1), (4.47, 4.1), (7.99, 4.1)]
    for (title, lines), (x, y) in zip(cards, positions):
        add_card(slide, x, y, 3.18, 1.82, title, lines, 16)
    add_footer(slide)


def slide_team(prs, bg_path: Path = BG_PATH):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg_path)
    add_title(slide, "团队", "以清晰组织分工支撑快速响应与稳定交付", section="01")

    nodes = [
        ("团队经理", 1.0, 1.75, 2.1, 0.62),
        ("招商主管", 0.95, 3.0, 2.15, 0.62),
        ("业务主管", 3.55, 3.0, 2.15, 0.62),
        ("业务员", 6.15, 2.35, 1.65, 0.58),
        ("跟单员", 6.15, 3.15, 1.65, 0.58),
        ("运营支持", 6.15, 3.95, 1.65, 0.58),
    ]
    for text, x, y, w, h in nodes:
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, emu(x), emu(y), emu(w), emu(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_FILL
        shape.line.color.rgb = CARD_LINE
        set_shape_text(shape, text, 14, WHITE, True)

    connectors = [
        ((2.05, 2.37), (2.03, 3.0)),
        ((2.9, 2.37), (4.63, 3.0)),
        ((5.7, 3.31), (6.15, 2.64)),
        ((5.7, 3.31), (6.15, 3.44)),
        ((5.7, 3.31), (6.15, 4.24)),
    ]
    for (x1, y1), (x2, y2) in connectors:
        line = slide.shapes.add_connector(MSO_CONNECTOR_TYPE.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2))
        line.line.color.rgb = ACCENT_2
        line.line.width = Pt(1.4)

    add_card(slide, 8.25, 1.78, 4.1, 3.55, "团队建设重点", [
        "组织架构清晰，管理层与执行层职责分明",
        "角色分工明确，避免职责交叉与执行空档",
        "协作机制统一，保证信息传递高效顺畅",
        "快速响应客户与订单需求，提升团队执行力",
    ], 17)
    highlight = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, emu(0.95), emu(5.45), emu(11.38), emu(0.9))
    highlight.fill.solid()
    highlight.fill.fore_color.rgb = RGBColor(11, 52, 111)
    highlight.fill.transparency = 0.15
    highlight.line.color.rgb = CARD_LINE
    tf = highlight.text_frame
    tf.text = "核心目标：以“清晰结构 + 高效协作 + 快速响应”为团队运行基础，支撑业务规模化发展。"
    tf.paragraphs[0].font.size = Pt(15)
    tf.paragraphs[0].font.color.rgb = TEXT
    tf.paragraphs[0].font.bold = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    add_footer(slide)


def slide_policy(prs, bg_path: Path = BG_PATH):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg_path)
    add_title(slide, "团队制度", "通过统一标准减少内耗，保障协同效率", section="02")
    items = [
        ("考勤与请假", "明确出勤要求、请假审批与异常反馈机制"),
        ("会议与汇报", "固定晨会/周会节奏，形成目标、进度、问题同步闭环"),
        ("资料管理", "统一客户资料、报价文件、订单记录的归档标准"),
        ("沟通规范", "统一内部协作、跨岗位交接与问题升级路径"),
    ]
    y = 1.7
    for title, desc in items:
        band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, emu(0.98), emu(y), emu(11.35), emu(0.95))
        band.fill.solid()
        band.fill.fore_color.rgb = CARD_FILL
        band.fill.transparency = 0.08
        band.line.color.rgb = CARD_LINE
        icon = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, emu(1.22), emu(y + 0.23), emu(0.32), emu(0.32))
        icon.fill.solid()
        icon.fill.fore_color.rgb = ACCENT
        icon.line.fill.background()
        t1 = slide.shapes.add_textbox(emu(1.7), emu(y + 0.15), emu(2.2), emu(0.28))
        p1 = t1.text_frame.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = WHITE
        t2 = slide.shapes.add_textbox(emu(4.0), emu(y + 0.15), emu(7.7), emu(0.45))
        p2 = t2.text_frame.paragraphs[0]
        p2.text = desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT
        y += 1.12
    add_card(slide, 0.98, 6.18, 11.35, 0.72, "制度建设导向", ["统一标准、强化执行、沉淀流程、保障协同。"], 16)
    add_footer(slide)


def slide_roles(prs, bg_path: Path = BG_PATH):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg_path)
    add_title(slide, "岗位职责", "责任到人、协同有序，确保每个岗位都能稳定产出", section="03")
    headers = [(1.05, 1.72, 2.1, "岗位"), (3.25, 1.72, 8.9, "核心职责")]
    for x, y, w, text in headers:
        h = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, emu(x), emu(y), emu(w), emu(0.56))
        h.fill.solid()
        h.fill.fore_color.rgb = RGBColor(15, 69, 136)
        h.line.color.rgb = CARD_LINE
        set_shape_text(h, text, 14, WHITE, True)
    rows = [
        ("业务员", "负责客户开发、询盘跟进、报价推进与成交转化"),
        ("跟单员", "负责订单执行、节点跟踪、资料对接与内部协调"),
        ("运营", "负责平台维护、资料整理、样品信息与数据支持"),
        ("主管", "负责团队管理、目标拆解、过程督导与问题协调"),
        ("经理", "负责策略制定、资源配置、结果复盘与机制优化"),
    ]
    y = 2.4
    for job, duty in rows:
        left = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, emu(1.05), emu(y), emu(2.1), emu(0.72))
        left.fill.solid()
        left.fill.fore_color.rgb = CARD_FILL
        left.line.color.rgb = CARD_LINE
        set_shape_text(left, job, 14, WHITE, True)
        right = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, emu(3.25), emu(y), emu(8.9), emu(0.72))
        right.fill.solid()
        right.fill.fore_color.rgb = RGBColor(9, 33, 73)
        right.fill.transparency = 0.08
        right.line.color.rgb = CARD_LINE
        right.text_frame.text = duty
        right.text_frame.paragraphs[0].font.size = Pt(13)
        right.text_frame.paragraphs[0].font.color.rgb = TEXT
        right.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        y += 0.84
    add_footer(slide)


def slide_workflow(prs, bg_path: Path = BG_PATH):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg_path)
    add_title(slide, "工作流程", "把日常执行动作标准化，确保任务推进可视、可控、可复盘", section="04")
    steps = ["接收任务", "客户信息", "跟进沟通", "提交报价", "内部协同", "结果反馈", "归档总结"]
    x = 0.72
    for idx, step in enumerate(steps):
        node = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, emu(x), emu(2.75), emu(1.6), emu(0.95))
        node.fill.solid()
        node.fill.fore_color.rgb = CARD_FILL
        node.line.color.rgb = ACCENT if idx in (0, len(steps)-1) else CARD_LINE
        num = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, emu(x + 0.1), emu(2.49), emu(0.36), emu(0.36))
        num.fill.solid()
        num.fill.fore_color.rgb = ACCENT
        num.line.fill.background()
        set_shape_text(num, str(idx + 1), 11, RGBColor(7, 29, 67), True)
        set_shape_text(node, step, 13, WHITE, True)
        if idx < len(steps) - 1:
            arrow = slide.shapes.add_connector(MSO_CONNECTOR_TYPE.STRAIGHT, emu(x + 1.6), emu(3.22), emu(x + 1.88), emu(3.22))
            arrow.line.color.rgb = ACCENT_2
            arrow.line.width = Pt(1.6)
        x += 1.78
    add_card(slide, 1.1, 4.45, 10.95, 1.35, "流程价值", [
        "统一执行步骤，减少重复沟通与流程遗漏；",
        "提高报价、跟进、反馈与归档的衔接效率。",
    ], 16)
    add_footer(slide)


def slide_business(prs, bg_path: Path = BG_PATH):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg_path)
    add_title(slide, "业务流程", "贯穿客户开发到售后复盘的完整成交链路", section="05")
    center_x, center_y = 6.65, 3.75
    radius_x, radius_y = 4.2, 1.95
    labels = ["获取询盘", "客户开发", "需求确认", "样品报价", "商务谈判", "签单成交", "订单执行", "售后复盘"]
    for idx, label in enumerate(labels):
        angle = -pi / 2 + idx * (2 * pi / len(labels))
        x = center_x + radius_x * cos(angle) - 0.68
        y = center_y + radius_y * sin(angle) - 0.3
        node = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, emu(x), emu(y), emu(1.36), emu(0.6))
        node.fill.solid()
        node.fill.fore_color.rgb = CARD_FILL
        node.line.color.rgb = ACCENT if idx % 2 == 0 else CARD_LINE
        tf = node.text_frame
        tf.text = label
        tf.paragraphs[0].font.size = Pt(12.5)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = WHITE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    core = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, emu(5.3), emu(3.02), emu(2.7), emu(1.45))
    core.fill.solid()
    core.fill.fore_color.rgb = RGBColor(10, 57, 115)
    core.line.color.rgb = ACCENT_2
    set_shape_text(core, "业务增长闭环\n标准化推进", 16, WHITE, True)
    add_footer(slide)


def slide_perf(prs, bg_path: Path = BG_PATH):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg_path)
    add_title(slide, "薪酬绩效", "以科学激励机制驱动目标达成与团队成长", section="06")
    add_card(slide, 0.98, 1.8, 5.3, 4.6, "薪酬结构", [
        "基础薪资：保障岗位稳定产出与团队基本配置",
        "业绩提成：与成交贡献直接挂钩，强化结果导向",
        "月度考核：聚焦过程执行、客户跟进与协同质量",
    ], 17)
    add_card(slide, 7.03, 1.8, 5.3, 4.6, "绩效机制", [
        "季度奖励：鼓励阶段突破与优秀表现复制",
        "年度激励：绑定长期目标，提升核心成员稳定性",
        "奖惩机制：明确标准，形成“有结果、有激励”的文化氛围",
    ], 17)
    bar_base = 5.15
    widths = [1.5, 2.4, 3.5]
    labels = ["保底", "绩效", "激励"]
    for idx, (w, lab) in enumerate(zip(widths, labels)):
        bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, emu(1.3), emu(bar_base + idx * 0.34), emu(w), emu(0.18))
        bar.fill.solid()
        bar.fill.fore_color.rgb = [RGBColor(46, 110, 197), RGBColor(35, 161, 255), RGBColor(111, 255, 233)][idx]
        bar.line.fill.background()
        txt = slide.shapes.add_textbox(emu(1.3 + w + 0.15), emu(bar_base + idx * 0.28 - 0.02), emu(1.0), emu(0.24))
        p = txt.text_frame.paragraphs[0]
        p.text = lab
        p.font.size = Pt(11)
        p.font.color.rgb = MUTED
    add_footer(slide)


def slide_closing(prs, bg_path: Path = BG_PATH):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg_path)
    ring = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, emu(9.3), emu(1.08), emu(2.2), emu(2.2))
    ring.fill.background()
    ring.line.color.rgb = ACCENT_2
    ring.line.width = Pt(2)
    title = slide.shapes.add_textbox(emu(1.0), emu(2.1), emu(8.6), emu(1.4))
    tf = title.text_frame
    p = tf.paragraphs[0]
    p.text = "泰卦外贸团队系统"
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p2 = tf.add_paragraph()
    p2.text = "让管理更高效，让协作更顺畅，让业务更有结果"
    p2.font.size = Pt(16)
    p2.font.color.rgb = MUTED
    p2.space_before = Pt(12)
    bottom = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, emu(1.0), emu(4.4), emu(11.3), emu(1.0))
    bottom.fill.solid()
    bottom.fill.fore_color.rgb = RGBColor(10, 41, 92)
    bottom.fill.transparency = 0.12
    bottom.line.color.rgb = CARD_LINE
    tf = bottom.text_frame
    tf.text = "通过组织建设、制度规范、流程标准化与绩效激励的系统协同，持续提升外贸团队专业化水平与市场竞争力。"
    tf.paragraphs[0].font.size = Pt(15)
    tf.paragraphs[0].font.color.rgb = TEXT
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    thanks = slide.shapes.add_textbox(emu(5.05), emu(6.15), emu(3.2), emu(0.35))
    p = thanks.text_frame.paragraphs[0]
    p.text = "感谢聆听"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_2
    p.alignment = PP_ALIGN.CENTER
    add_footer(slide)



def build_presentation(output_path: Optional[Path] = None, background_path: Optional[Path] = None):
    """Create the shared background asset if needed and write the PPTX deck to the resolved target path."""
    background = Path(background_path) if background_path else BG_PATH
    if not background.is_absolute():
        background = ROOT / background
    ensure_background(background)
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide_cover(prs, background)
    slide_overview(prs, background)
    slide_team(prs, background)
    slide_policy(prs, background)
    slide_roles(prs, background)
    slide_workflow(prs, background)
    slide_business(prs, background)
    slide_perf(prs, background)
    slide_closing(prs, background)
    target = Path(output_path) if output_path else OUTPUT
    if not target.is_absolute():
        target = ROOT / target
    target.parent.mkdir(parents=True, exist_ok=True)
    prs.save(target)
    return target


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate the 泰卦外贸团队系统 PowerPoint deck.")
    parser.add_argument("--output", default=str(OUTPUT), help="Optional output PPTX path")
    parser.add_argument("--background", default=str(BG_PATH), help="Optional background asset path")
    args = parser.parse_args()
    output = build_presentation(Path(args.output), Path(args.background))
    print(f"Created {output}")
